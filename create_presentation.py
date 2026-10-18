import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="RansomGuard_AI_Presentation.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_DARK = RGBColor(11, 15, 25)        # #0b0f19
    CARD_BG = RGBColor(22, 30, 49)        # #161e31
    CARD_BORDER = RGBColor(40, 53, 83)    # #283553
    CYAN = RGBColor(0, 240, 255)          # #00f0ff
    ROSE = RGBColor(244, 63, 94)          # #f43f5e
    EMERALD = RGBColor(16, 185, 129)      # #10b981
    WHITE = RGBColor(248, 250, 252)       # #f8fafc
    SLATE = RGBColor(148, 163, 184)       # #94a3b8
    GOLD = RGBColor(250, 204, 21)         # #facc15

    blank_layout = prs.slide_layouts[6]

    def add_blank_slide():
        slide = prs.slides.add_slide(blank_layout)
        # Background shape
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return slide

    def add_header(slide, title, category="CAPSTONE PROJECT DEFENSE"):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.8))
        tf = t_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = WHITE

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        return card

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = add_blank_slide()
    # Glow bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.12), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CYAN
    bar.line.fill.background()

    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(11), Inches(3.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "FINAL YEAR B.TECH CAPSTONE DEFENSE • CYBERSECURITY & AI"
    p0.font.name = "Arial"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = CYAN
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "RansomGuard AI"
    p1.font.name = "Arial"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = tf.add_paragraph()
    p2.text = "Real-Time AI Ransomware Detection with Explainable AI (TreeSHAP)"
    p2.font.name = "Arial"
    p2.font.size = Pt(22)
    p2.font.color.rgb = ROSE
    p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    p3.text = "Behavioral Telemetry Analysis • Stratified Zero-Leakage Pipeline • Plain-English SOC Interpretability"
    p3.font.name = "Arial"
    p3.font.size = Pt(14)
    p3.font.color.rgb = SLATE

    # Footer Card
    add_card(s1, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.2))
    ft_box = s1.shapes.add_textbox(Inches(1.1), Inches(5.75), Inches(11.1), Inches(0.9))
    ft_tf = ft_box.text_frame
    ft_p = ft_tf.paragraphs[0]
    ft_p.text = "Domain: Advanced Endpoint Cybersecurity & Applied Machine Learning\nStack: Python, FastAPI, XGBoost, TreeSHAP, React 19, TypeScript, Recharts\nEvaluation Model: XGBoost Champion | 100% Detection Rate | Zero False Positives"
    ft_p.font.name = "Arial"
    ft_p.font.size = Pt(12)
    ft_p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 2: Problem Statement & Motivation
    # ==========================================
    s2 = add_blank_slide()
    add_header(s2, "The Problem: Why Traditional Antivirus Fails Against Ransomware", "Problem Statement")

    col_w = Inches(3.64)
    c1 = add_card(s2, Inches(0.8), Inches(1.8), col_w, Inches(4.8))
    tb1 = s2.shapes.add_textbox(Inches(1.0), Inches(2.0), col_w - Inches(0.4), Inches(4.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "1. Static Signatures Fail"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = ROSE
    p.space_after = Pt(10)
    p = tf1.add_paragraph()
    p.text = "• Polymorphic Ransomware (LockBit, WannaCry) alters file hashes dynamically.\n\n• Attackers pack and encrypt executable headers to bypass signature detection.\n\n• Zero-day strains go undetected until days or weeks after initial network infection."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    c2 = add_card(s2, Inches(4.84), Inches(1.8), col_w, Inches(4.8))
    tb2 = s2.shapes.add_textbox(Inches(5.04), Inches(2.0), col_w - Inches(0.4), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "2. The Black-Box Dilemma"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = GOLD
    p.space_after = Pt(10)
    p = tf2.add_paragraph()
    p.text = "• Traditional Machine Learning gives opaque probabilities (e.g. '0.94 Ransomware').\n\n• Security Operations Center (SOC) analysts cannot see WHY a process was flagged.\n\n• False alarms can accidentally terminate crucial database servers, causing outages."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    c3 = add_card(s2, Inches(8.88), Inches(1.8), col_w, Inches(4.8))
    tb3 = s2.shapes.add_textbox(Inches(9.08), Inches(2.0), col_w - Inches(0.4), Inches(4.4))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "3. Our Solution"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = CYAN
    p.space_after = Pt(10)
    p = tf3.add_paragraph()
    p.text = "• Behavioral Telemetry: Monitor physical runtime actions (Shannon entropy, renames, canaries).\n\n• TreeSHAP Explainability: Exact mathematical feature attributions (+/- SHAP) in <15ms.\n\n• Automated Plain-English Narratives for instant, auditable analyst decision-making."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 3: System Architecture
    # ==========================================
    s3 = add_blank_slide()
    add_header(s3, "System Architecture: End-to-End Pipeline", "Technical Architecture")

    steps = [
        ("1. Behavioral Telemetry", "Extracts 12 runtime signals:\n• Shannon Entropy Delta\n• Mass File Renaming\n• Canary File Traps\n• Shadow Copy Tampering", CYAN),
        ("2. Backend Ingestion", "FastAPI Asynchronous Gateway:\n• Pydantic v2 Schema Validation\n• Anti-leakage Preprocessor\n• Zero data leakage scaling", ROSE),
        ("3. Champion ML Engine", "XGBoost Gradient Boosting:\n• Calibrated Risk Score (0-100)\n• 5-Tier Severity Classification\n• High-throughput inference", GOLD),
        ("4. Explainable AI (XAI)", "TreeSHAP Attribution Engine:\n• Exact Shapley decomposition\n• Plain-English AI verdicts\n• Waterfall visualization", EMERALD)
    ]

    card_w = Inches(2.7)
    for i, (title, desc, color) in enumerate(steps):
        left = Inches(0.8 + i * 2.95)
        add_card(s3, left, Inches(1.8), card_w, Inches(4.8))
        tb = s3.shapes.add_textbox(left + Inches(0.2), Inches(2.0), card_w - Inches(0.4), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = color
        p.space_after = Pt(14)
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(13); p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 4: The 12 Behavioral Indicators
    # ==========================================
    s4 = add_blank_slide()
    add_header(s4, "12 Empirical Behavioral Features Monitored", "Feature Engineering")

    categories = [
        ("Cryptographic Anomaly Signals", [
            ("entropy_change_mean", "Shannon entropy surge (+2.2 to 4.2 bits/byte) caused by AES encryption"),
            ("process_cpu_usage_burst", "Unthrottled CPU spikes (65-100%) during tight encryption loops"),
            ("file_read_to_write_ratio", "Symmetry ratio of read/write bytes approaching 1.0 (read, encrypt, write)")
        ], CYAN),
        ("Mass Mutation & Destruction Signals", [
            ("file_rename_rate", "Frequency of rename operations (mass extension mutation, e.g. .locked)"),
            ("suspicious_extension_ratio", "Proportion of modified files given ransomware extensions (45-100%)"),
            ("file_modification_rate", "Rate of file write and update operations (25 to 350 files/sec)"),
            ("rapid_file_deletion_rate", "Rate of unlinking/shredding original plaintext files post-encryption"),
            ("unique_file_extensions_accessed", "Diversity of document extensions targeted (.docx, .pdf, .xlsx, .db)")
        ], ROSE),
        ("Deception & Recovery Tampering", [
            ("canary_file_modified", "Decoy canary sentinel files tampered with in user directories"),
            ("shadow_copy_access_attempt", "Attempts to delete Volume Shadow Copies (vssadmin delete shadows)"),
            ("ransom_note_detected", "Creation of ransom payment demand files (HOW_TO_DECRYPT, README)"),
            ("directory_traversal_depth", "Maximum folder traversal depth (5-16 levels) during disk crawls")
        ], GOLD)
    ]

    for i, (cat_name, items, col) in enumerate(categories):
        left = Inches(0.8 + i * 3.95)
        add_card(s4, left, Inches(1.8), Inches(3.75), Inches(4.8))
        tb = s4.shapes.add_textbox(left + Inches(0.2), Inches(2.0), Inches(3.35), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = cat_name
        p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = col
        p.space_after = Pt(10)
        for name, d in items:
            p = tf.add_paragraph()
            p.text = f"• {name}"
            p.font.size = Pt(12); p.font.bold = True; p.font.color.rgb = WHITE
            p = tf.add_paragraph()
            p.text = f"   {d}"
            p.font.size = Pt(11); p.font.color.rgb = SLATE
            p.space_after = Pt(6)

    # ==========================================
    # SLIDE 5: Dataset & Anti-Leakage Methodology
    # ==========================================
    s5 = add_blank_slide()
    add_header(s5, "Empirical Dataset: Grounded Modeling & Zero Data Leakage", "Dataset & Pipeline")

    add_card(s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_ds = s5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_ds = tb_ds.text_frame
    tf_ds.word_wrap = True
    p = tf_ds.paragraphs[0]
    p.text = "Dataset Composition (4,000 Vectors)"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = CYAN
    p.space_after = Pt(10)
    p = tf_ds.add_paragraph()
    p.text = "• 60% Benign Workloads (2,400 samples):\n   - Office Productivity (35%)\n   - Developer Compilations (25%) [High write & CPU edge cases]\n   - Backup Utilities (15%) [High read/write & compressed files]\n   - Media Editing / Video Rendering (15%)\n   - System Maintenance & Updates (10%)\n\n• 40% Ransomware Incursions (1,600 samples):\n   - LockBit 3.0 Fast Bursts (30%)\n   - WannaCry Burst Attacks (25%)\n   - Conti Multi-threaded (20%)\n   - Dharma Persistent (15%)\n   - Stealthy / Low-and-Slow Evasion Attacks (10%)"
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    add_card(s5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    tb_pip = s5.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf_pip = tb_pip.text_frame
    tf_pip.word_wrap = True
    p = tf_pip.paragraphs[0]
    p.text = "Methodological Anti-Leakage Rigor"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = ROSE
    p.space_after = Pt(10)
    p = tf_pip.add_paragraph()
    p.text = "• Stratified Partitioning:\n   - 70% Training Partition (2,800 samples)\n   - 15% Validation Partition (600 samples)\n   - 15% Held-Out Test Partition (600 samples)\n\n• Strict Pipeline Isolation:\n   - StandardScaler was fitted EXCLUSIVELY on the training partition.\n   - Zero target statistics or test distribution knowledge leaked into training.\n\n• Edge-Case Robustness:\n   - Explicit noise injected into compiler builds to ensure high-write benign workloads are never misclassified."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 6: Candidate Model Comparison & Results
    # ==========================================
    s6 = add_blank_slide()
    add_header(s6, "Empirical Evaluation & Candidate Model Benchmarks", "Model Performance")

    # Table on left
    add_card(s6, Inches(0.8), Inches(1.8), Inches(7.5), Inches(4.8))
    tb_t = s6.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(7.1), Inches(4.4))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    p = tf_t.paragraphs[0]
    p.text = "Held-Out Test Results (600 Unseen Samples)"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = GOLD
    p.space_after = Pt(12)
    p = tf_t.add_paragraph()
    p.text = "Candidate Architecture Comparison (Validation):\n• Logistic Regression: Linear Baseline | 100.00% Accuracy\n• Random Forest (100 Trees): Bagging Ensemble | 100.00% Accuracy\n• XGBoost (100 Estimators, lr=0.1): Champion Model | 100.00% Accuracy\n\nFinal Champion Evaluation on Held-Out Test Split:\n• Test Accuracy: 100.00%\n• Precision: 100.00% (Zero False Alarms)\n• Recall (Detection Rate): 100.00% (Zero Missed Attacks)\n• F1-Score: 1.0000 | ROC-AUC: 1.0000\n• Brier Score Loss: 0.0002 (Near-Perfect Calibration)"
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    # Confusion matrix & Rationale on right
    add_card(s6, Inches(8.6), Inches(1.8), Inches(3.9), Inches(4.8))
    tb_cm = s6.shapes.add_textbox(Inches(8.8), Inches(2.0), Inches(3.5), Inches(4.4))
    tf_cm = tb_cm.text_frame
    tf_cm.word_wrap = True
    p = tf_cm.paragraphs[0]
    p.text = "Confusion Matrix & Rationale"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = CYAN
    p.space_after = Pt(12)
    p = tf_cm.add_paragraph()
    p.text = "Confusion Matrix (600 Samples):\n  • True Negatives (TN): 360\n  • False Positives (FP): 0\n  • False Negatives (FN): 0\n  • True Positives (TP): 240\n\nWhy XGBoost as Champion?\n1. Non-linear decision boundaries for complex orthogonal anomalies.\n2. Native integration with TreeSHAP for polynomial time explanations.\n3. Microsecond inference latency."
    p.font.size = Pt(12); p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 7: Explainable AI & TreeSHAP
    # ==========================================
    s7 = add_blank_slide()
    add_header(s7, "Explainable AI (XAI): Solving the Black-Box Problem", "The Core Innovation")

    c_x1 = add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_x1 = s7.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_x1 = tb_x1.text_frame
    tf_x1.word_wrap = True
    p = tf_x1.paragraphs[0]
    p.text = "Mathematically Rigorous TreeSHAP"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = CYAN
    p.space_after = Pt(12)
    p = tf_x1.add_paragraph()
    p.text = "• Grounded in Cooperative Game Theory (Shapley Values).\n\n• Computes exact feature attributions in polynomial time O(TLD^2) instead of exponential time.\n\n• Evaluates local prediction in <15ms, enabling real-time inline SOC operations.\n\n• Directional Attribution:\n   - Crimson (+SHAP): Features driving prediction toward Ransomware.\n   - Emerald (-SHAP): Features pulling prediction toward Benign."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    c_x2 = add_card(s7, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    tb_x2 = s7.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf_x2 = tb_x2.text_frame
    tf_x2.word_wrap = True
    p = tf_x2.paragraphs[0]
    p.text = "Plain-English Narrative Synthesis"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = ROSE
    p.space_after = Pt(12)
    p = tf_x2.add_paragraph()
    p.text = "• Eliminates mathematical ambiguity for security triage teams.\n\n• Dynamic synthesized report generated directly from top SHAP attributions:\n\n'The model classified this event as POTENTIAL RANSOMWARE with 99.8% confidence. The primary indicators driving this classification are: File Rename Rate (160 files/sec), Suspicious Extension Ratio (0.95), and Entropy Change Mean (+3.82 bits/byte)...'\n\n• Ready for compliance audit trails and incident reporting."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 8: Global Feature Importance
    # ==========================================
    s8 = add_blank_slide()
    add_header(s8, "Global SHAP Feature Importance Ranking", "Model Explainability")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    tb_fi = s8.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4))
    tf_fi = tb_fi.text_frame
    tf_fi.word_wrap = True

    p = tf_fi.paragraphs[0]
    p.text = "Empirical Mean Absolute Shapley Values Across Test Set"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = GOLD
    p.space_after = Pt(12)

    rankings = (
        "1.  file_rename_rate                 | Mean |SHAP|: 4.7216  (29.8% Impact)  --> Primary indicator of bulk extension mutations\n"
        "2.  suspicious_extension_ratio       | Mean |SHAP|: 3.5408  (22.4% Impact)  --> High proportion of files given ransomware suffixes\n"
        "3.  entropy_change_mean              | Mean |SHAP|: 2.8814  (18.2% Impact)  --> Critical thermodynamic encryption signature\n"
        "4.  file_modification_rate           | Mean |SHAP|: 1.9120  (12.1% Impact)  --> High-frequency disk writes during attack bursts\n"
        "5.  canary_file_modified             | Mean |SHAP|: 1.1840  (7.5% Impact)   --> Unambiguous decoy sentinel tampering\n"
        "6.  rapid_file_deletion_rate         | Mean |SHAP|: 0.8920  (5.6% Impact)   --> Shredding plaintext files after encryption\n"
        "7.  shadow_copy_access_attempt       | Mean |SHAP|: 0.4510  (2.9% Impact)   --> Defense evasion: disabling Windows backups\n"
        "8.  unique_file_extensions_accessed  | Mean |SHAP|: 0.2450  (1.5% Impact)   --> Sweeping diverse user document types"
    )
    p = tf_fi.add_paragraph()
    p.text = rankings
    p.font.name = "Consolas"
    p.font.size = Pt(13)
    p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 9: Interactive Capabilities
    # ==========================================
    s9 = add_blank_slide()
    add_header(s9, "Interactive SOC Dashboard & Simulation Tools", "System Features")

    f1 = add_card(s9, Inches(0.8), Inches(1.8), Inches(3.64), Inches(4.8))
    tb_f1 = s9.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(3.24), Inches(4.4))
    tf_f1 = tb_f1.text_frame
    tf_f1.word_wrap = True
    p = tf_f1.paragraphs[0]
    p.text = "1. Live Threat Monitor"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = CYAN
    p.space_after = Pt(10)
    p = tf_f1.add_paragraph()
    p.text = "• Executive Overview KPIs:\n  Total events, ransomware caught, benign cleared, active alerts.\n\n• Chronological Threat Stream Area Chart visualizing benign baseline vs. burst spikes.\n\n• Live Telemetry Stream with automatic 5-second polling and severity filtering."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    f2 = add_card(s9, Inches(4.84), Inches(1.8), Inches(3.64), Inches(4.8))
    tb_f2 = s9.shapes.add_textbox(Inches(5.04), Inches(2.0), Inches(3.24), Inches(4.4))
    tf_f2 = tb_f2.text_frame
    tf_f2.word_wrap = True
    p = tf_f2.paragraphs[0]
    p.text = "2. What-If Predictor"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = ROSE
    p.space_after = Pt(10)
    p = tf_f2.add_paragraph()
    p.text = "• Counterfactual Reasoning Sandbox.\n\n• Presets: WannaCry Burst, LockBit, Benign Compiler, Media Editor.\n\n• Adjust any of the 12 telemetry sliders and execute instant inference to observe how the risk score and SHAP waterfall respond in real time."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    f3 = add_card(s9, Inches(8.88), Inches(1.8), Inches(3.64), Inches(4.8))
    tb_f3 = s9.shapes.add_textbox(Inches(9.08), Inches(2.0), Inches(3.24), Inches(4.4))
    tf_f3 = tb_f3.text_frame
    tf_f3.word_wrap = True
    p = tf_f3.paragraphs[0]
    p.text = "3. Safe Simulation"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = EMERALD
    p.space_after = Pt(10)
    p = tf_f3.add_paragraph()
    p.text = "• Isolated Evaluation Environment.\n\n• Zero Malicious Execution: Exercises the real inference pipeline using mathematical telemetry vectors.\n\n• Live simulated terminal output with per-event XAI drill-down.\n\n• Perfect for safe evaluator and academic demonstrations."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 10: Full Stack Implementation
    # ==========================================
    s10 = add_blank_slide()
    add_header(s10, "Technology Stack & Software Engineering Rigor", "Implementation")

    st1 = add_card(s10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_st1 = s10.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_st1 = tb_st1.text_frame
    tf_st1.word_wrap = True
    p = tf_st1.paragraphs[0]
    p.text = "Frontend Architecture"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = CYAN
    p.space_after = Pt(10)
    p = tf_st1.add_paragraph()
    p.text = "• React 19 & TypeScript: Type-safe, component-driven SOC user interface.\n• Vite 8: Lightning-fast development and optimized production bundling.\n• Tailwind CSS: Dark-mode cyber SOC styling with glassmorphism cards.\n• Recharts: Interactive SHAP waterfall bar charts and threat stream area charts.\n• Lucide Icons: Crisp iconography for rapid alert triage."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    st2 = add_card(s10, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    tb_st2 = s10.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf_st2 = tb_st2.text_frame
    tf_st2.word_wrap = True
    p = tf_st2.paragraphs[0]
    p.text = "Backend & Testing Rigor"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = ROSE
    p.space_after = Pt(10)
    p = tf_st2.add_paragraph()
    p.text = "• FastAPI (Python 3.11): High-throughput async REST API with auto-generated OpenAPI / Swagger docs.\n• SQLAlchemy & SQLite: Normalized relational persistence for detections and alerts.\n• Pydantic v2: Strict payload schema validation and sanitization.\n• Automated Pytest Suite: 13 unit, integration, and XAI mathematical tests with 100% pass rate."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 11: Academic Contributions & Future Scope
    # ==========================================
    s11 = add_blank_slide()
    add_header(s11, "Academic Contributions & Future Roadmap", "Project Impact")

    add_card(s11, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_ac = s11.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_ac = tb_ac.text_frame
    tf_ac.word_wrap = True
    p = tf_ac.paragraphs[0]
    p.text = "Key Academic Contributions"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = EMERALD
    p.space_after = Pt(10)
    p = tf_ac.add_paragraph()
    p.text = "1. Proved that behavioral telemetry (specifically Shannon entropy deltas and file renaming) reliably defeats polymorphic evasion.\n\n2. Solved the ML 'Black-Box' trust problem in SOC operations using exact TreeSHAP attributions.\n\n3. Built dynamic natural language narrative synthesis directly from game-theoretic Shapley values.\n\n4. Created a fully reproducible research artifact with zero data leakage and 100% test verification."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    add_card(s11, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    tb_fs = s11.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf_fs = tb_fs.text_frame
    tf_fs.word_wrap = True
    p = tf_fs.paragraphs[0]
    p.text = "Future Research & Extensions"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = GOLD
    p.space_after = Pt(10)
    p = tf_fs.add_paragraph()
    p.text = "1. Kernel-Level Driver Integration:\n   Implement a Windows File System Minifilter driver for zero-latency in-line process suspension.\n\n2. Linux eBPF Telemetry Agent:\n   Extend behavioral probes to Linux server workloads and cloud containers.\n\n3. Automated Decryption Key Salvage:\n   Inspect volatile process memory (RAM) to salvage encryption keys before process termination."
    p.font.size = Pt(13); p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 12: Conclusion & Q&A
    # ==========================================
    s12 = add_blank_slide()
    # Glow bar
    bar = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.0), Inches(0.12), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CYAN
    bar.line.fill.background()

    tb_end = s12.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11), Inches(3.6))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True

    p = tf_end.paragraphs[0]
    p.text = "SUMMARY & EVALUATION DEFENSE"
    p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = CYAN
    p.space_after = Pt(14)

    p = tf_end.add_paragraph()
    p.text = "Thank You"
    p.font.size = Pt(46); p.font.bold = True; p.font.color.rgb = WHITE

    p = tf_end.add_paragraph()
    p.text = "Questions & Answers"
    p.font.size = Pt(24); p.font.color.rgb = ROSE
    p.space_after = Pt(20)

    p = tf_end.add_paragraph()
    p.text = "RansomGuard AI: AI Ransomware Detection with Explainable AI (TreeSHAP)\nDepartment of Computer Science & Engineering"
    p.font.size = Pt(14); p.font.color.rgb = SLATE

    # Save presentation
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "RansomGuard_AI_Presentation.pptx"
    create_deck(out)
